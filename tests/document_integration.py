"""Real LibreOffice/Poppler integration; fixture OCR deliberately uses mock models."""
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfgen.canvas import Canvas


def run(binary):
    assert shutil.which('soffice') and shutil.which('pdftoppm') and shutil.which('pdfinfo')
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        doc = Document()
        doc.add_paragraph('OmniOCR Word fixture')
        doc.save(root / "test 'quote'.docx")
        pres = Presentation()
        slide = pres.slides.add_slide(pres.slide_layouts[0])
        slide.shapes.title.text = 'OmniOCR PPT fixture'
        pres.save(root / 'test.pptx')
        book = Workbook()
        book.active['A1'] = 'OmniOCR Excel fixture'
        book.save(root / 'test.xlsx')
        pdf = Canvas(str(root / 'test.pdf'))
        for i in range(12):
            pdf.drawString(72, 720, f'Page {i+1}')
            pdf.showPage()
        pdf.save()
        small = Canvas(str(root / 'small.pdf'), pagesize=(72, 72))
        small.drawString(5, 35, 'small')
        small.showPage()
        small.save()
        # Exercise legacy binary formats with actual LibreOffice conversions too.
        legacy = []
        for source, target in [("test 'quote'.docx", 'doc'), ('test.pptx', 'ppt'), ('test.xlsx', 'xls')]:
            subprocess.run(['soffice', '-env:UserInstallation=' + (root / 'conversion-profile').as_uri(),
                            '--headless', '--convert-to', target, '--outdir', str(root), str(root / source)],
                           check=True, capture_output=True, timeout=60)
            legacy.append((root / source).with_suffix('.' + target).name)
        config = json.loads((Path(__file__).resolve().parents[1] / 'configs/demo.json').read_text())
        config['document']['max_pixels'] = 250000
        (root / 'config.json').write_text(json.dumps(config))
        for i, filename in enumerate(["test 'quote'.docx", 'test.pptx', 'test.xlsx', 'test.pdf', *legacy]):
            subprocess.run([binary, '--config', str(root / 'config.json'), '--input', str(root / filename),
                            '--output', str(root / f'out-{i}')], check=True, timeout=120)
            result = json.loads((root / f'out-{i}/result.json').read_text())
            expected = 12 if filename.endswith('.pdf') else 1
            assert [p['page'] for p in result['pages']] == list(range(1, expected + 1)), filename
        config['document']['dpi'] = 72
        (root / 'config.json').write_text(json.dumps(config))
        subprocess.run([binary, '--config', str(root / 'config.json'), '--input', str(root / 'small.pdf'),
                        '--output', str(root / 'small-out')], check=True, timeout=30)
        small_page = json.loads((root / 'small-out/result.json').read_text())['pages'][0]
        assert small_page['width'] == 72 and small_page['height'] == 72, 'max_pixels must not upscale small pages'
        config['document']['max_pages'] = 3
        (root / 'config.json').write_text(json.dumps(config))
        failed = subprocess.run([binary, '--config', str(root / 'config.json'), '--input', str(root / 'test.pdf'),
                                 '--output', str(root / 'limited')], capture_output=True, timeout=30)
        assert failed.returncode == 1 and b'page count exceeds limit' in failed.stderr
        # A zero-exit converter with no output must not be accepted.
        config['document']['soffice'] = '/bin/true'
        (root / 'config.json').write_text(json.dumps(config))
        failed = subprocess.run([binary, '--config', str(root / 'config.json'), '--input', str(root / 'test.xlsx'),
                                 '--output', str(root / 'no-pdf')], capture_output=True, timeout=30)
        assert failed.returncode == 1 and b'produced no PDF' in failed.stderr
    print('PASS: DOCX/PPTX/XLSX/DOC/PPT/XLS/PDF, 12-page order, page limit, converter false-success')


if __name__ == '__main__':
    run(str(Path(sys.argv[1]).resolve()))
