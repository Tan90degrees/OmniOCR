"""Actual new-format rendering with mock layout; no claim about OCR accuracy.

Default: RTF, ODF, HTML, CSV and multi-page TIFF. --external requires real
Calibre and the built OFDRW converter; it never silently skips missing tools.
"""
import argparse
import json
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import zipfile
from PIL import Image
from docx import Document
from pptx import Presentation
from openpyxl import Workbook


def epub_fixture(path):
    with zipfile.ZipFile(path, 'w') as book:
        book.writestr('mimetype', 'application/epub+zip', compress_type=zipfile.ZIP_STORED)
        book.writestr('META-INF/container.xml', '''<?xml version="1.0"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>''')
        book.writestr('OEBPS/content.opf', '''<?xml version="1.0"?>
<package xmlns="http://www.idpf.org/2007/opf" version="2.0" unique-identifier="bookid">
<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="bookid">omniocr-fixture</dc:identifier>
<dc:title>OmniOCR EPUB fixture</dc:title><dc:language>en</dc:language></metadata>
<manifest><item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/>
<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/></manifest>
<spine toc="ncx"><itemref idref="chapter"/></spine></package>''')
        book.writestr('OEBPS/chapter.xhtml', '''<html xmlns="http://www.w3.org/1999/xhtml">
<head><title>Chapter</title></head><body><h1>EPUB fixture</h1><p>FORMAT_FIXTURE</p></body></html>''')
        book.writestr('OEBPS/toc.ncx', '''<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1">
<head><meta name="dtb:uid" content="omniocr-fixture"/></head><docTitle><text>Fixture</text></docTitle>
<navMap><navPoint id="chapter" playOrder="1"><navLabel><text>Chapter</text></navLabel>
<content src="chapter.xhtml"/></navPoint></navMap></ncx>''')


def ofd_fixture(path):
    # Two vector-only pages exercise real OFD parsing, page order and rendering
    # without depending on a particular installed Chinese font.
    ns = 'http://www.ofdspec.org/2016'
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as doc:
        doc.writestr('OFD.xml', f'''<?xml version="1.0" encoding="UTF-8"?>
<ofd:OFD xmlns:ofd="{ns}" Version="1.0" DocType="OFD"><ofd:DocBody><ofd:DocInfo>
<ofd:DocID>0123456789abcdef0123456789abcdef</ofd:DocID></ofd:DocInfo>
<ofd:DocRoot>Doc_0/Document.xml</ofd:DocRoot></ofd:DocBody></ofd:OFD>''')
        doc.writestr('Doc_0/Document.xml', f'''<ofd:Document xmlns:ofd="{ns}">
<ofd:CommonData><ofd:MaxUnitID>20</ofd:MaxUnitID><ofd:PageArea>
<ofd:PhysicalBox>0 0 80 60</ofd:PhysicalBox></ofd:PageArea></ofd:CommonData>
<ofd:Pages><ofd:Page ID="1" BaseLoc="Pages/Page_0/Content.xml"/>
<ofd:Page ID="2" BaseLoc="Pages/Page_1/Content.xml"/></ofd:Pages></ofd:Document>''')
        for page, color in enumerate(('255 0 0', '0 0 255')):
            doc.writestr(f'Doc_0/Pages/Page_{page}/Content.xml', f'''<ofd:Page xmlns:ofd="{ns}">
<ofd:Content><ofd:Layer ID="{10 + page}" Type="Body"><ofd:PathObject ID="{15 + page}"
Boundary="10 10 50 30" Stroke="false" Fill="true"><ofd:FillColor Value="{color}"/>
<ofd:AbbreviatedData>M 0 0 L 50 0 L 50 30 L 0 30 C</ofd:AbbreviatedData>
</ofd:PathObject></ofd:Layer></ofd:Content></ofd:Page>''')


def run(binary, external=False, ofd_converter=None):
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        doc = Document(); doc.add_paragraph('ODT FORMAT_FIXTURE'); doc.save(root / 'word.docx')
        pres = Presentation(); pres.slides.add_slide(pres.slide_layouts[0]).shapes.title.text = 'ODP FORMAT_FIXTURE'
        pres.save(root / 'slides.pptx')
        wb = Workbook(); wb.active['A1'] = 'ODS FORMAT_FIXTURE'; wb.save(root / 'sheet.xlsx')
        fixtures = []
        for source, ext in [('word.docx', 'odt'), ('slides.pptx', 'odp'), ('sheet.xlsx', 'ods')]:
            subprocess.run(['soffice', '-env:UserInstallation=' + (root/'profile').as_uri(),
                            '--headless', '--convert-to', ext, '--outdir', str(root), str(root/source)],
                           capture_output=True, check=True, timeout=90)
            fixture = (root/source).with_suffix('.'+ext)
            assert fixture.is_file(), fixture
            fixtures.append(fixture)
        (root/'rich.RTF').write_text(r'{\rtf1\ansi\deff0 {\fonttbl {\f0 Helvetica;}}\f0\fs28 RTF FORMAT_FIXTURE\par}')
        # Relative HTML images must survive conversion from the original directory.
        Image.new('RGB', (40, 20), 'red').save(root/'relative.png')
        for ext in ('html', 'htm'):
            p = root / ("page 'quoted'."+ext)
            p.write_text('<!DOCTYPE html><html><head><meta charset="utf-8"></head>'
                         '<body><h1>HTML FORMAT_FIXTURE</h1><img src="relative.png"></body></html>')
            fixtures.append(p)
        (root/'cells.csv').write_text('\ufeffID,Text,Formula\r\n00123,"中文,quoted\nline",=1+2\r\n')
        fixtures += [root/'rich.RTF', root/'cells.csv']
        frames = [Image.new('RGB', (64, 32), color) for color in ('red', 'blue')]
        for ext in ('tif', 'tiff'):
            p = root / ('scan.'+ext)
            frames[0].save(p, save_all=True, append_images=frames[1:], compression='tiff_lzw')
            fixtures.append(p)
        if external:
            assert shutil.which('ebook-convert'), '--external requires Calibre ebook-convert'
            assert ofd_converter and Path(ofd_converter).is_file(), '--external requires built OFDRW wrapper'
            epub_fixture(root/'book.epub'); ofd_fixture(root/'pages.ofd')
            fixtures += [root/'book.epub', root/'pages.ofd']
        config = {'version': 1, 'document': {'dpi': 72, 'max_pixels': 500000, 'max_pages': 20},
                  'layout': {'provider': 'normalized', 'coordinates': 'normalized', 'model': 'layout'},
                  'models': {'layout': {'backend': 'mock', 'response': {'boxes': [{'type': 'image', 'bbox': [0, 0, 1, 1]}]}}},
                  'routes': {'*': {'action': 'image'}}}
        if ofd_converter:
            config['document']['ofd_converter'] = str(Path(ofd_converter).resolve())
        cfg = root/'config.json'
        def invoke(path, name, expected=0):
            cfg.write_text(json.dumps(config))
            out = root/name
            p = subprocess.run([binary, '--config', str(cfg), '--input', str(path), '--output', str(out)],
                               capture_output=True, text=True, timeout=150,
                               env={**os.environ, 'QT_QPA_PLATFORM': 'offscreen', 'QTWEBENGINE_DISABLE_SANDBOX': '1'})
            assert p.returncode == expected, (path, p.returncode, p.stderr)
            return out, p
        for i, path in enumerate(fixtures):
            out, _ = invoke(path, f'out-{i}')
            result = json.loads((out/'result.json').read_text())
            pages = result['pages']
            assert pages and [p['page'] for p in pages] == list(range(1, len(pages)+1))
            assert (out/'result.md').is_file()
            images = [Image.open(out/p['blocks'][0]['asset']).convert('RGB') for p in pages]
            assert any(im.getextrema() != ((255,255),(255,255),(255,255)) for im in images), path
            if path.suffix.lower() in ('.tif', '.tiff', '.ofd'):
                assert len(pages) == 2, path
                a, b = [im.getpixel((im.width//2, im.height//2)) for im in images]
                assert a[0] > a[2] and b[2] > b[0], (path, a, b)
            if path.suffix in ('.html', '.htm'):
                assert any(r > 200 and g < 80 and b < 80 for im in images for r,g,b in im.getdata()), 'relative HTML image missing'
            print('PASS:', path.suffix, len(pages), 'pages', flush=True)
        config['document']['max_pages'] = 1
        _, p = invoke(root/'scan.tiff', 'too-many-pages', 1)
        assert 'page count exceeds' in p.stderr
        config['document']['max_pages'] = 20; config['document']['max_pixels'] = 100
        _, p = invoke(root/'scan.tif', 'too-many-pixels', 1)
        assert 'max_pixels' in p.stderr
        config['document']['max_pixels'] = 500000
        for ext, key in (('epub','ebook_convert'), ('ofd','ofd_converter')):
            fake = root/('fake.'+ext); fake.write_bytes(b'not a valid document')
            config['document'][key] = '/bin/true'
            _, p = invoke(fake, 'no-output-'+ext, 1)
            assert 'produced no PDF' in p.stderr
            config['document'][key] = '/missing/converter'
            _, p = invoke(fake, 'missing-'+ext, 1)
            assert key in p.stderr and 'cannot start' in p.stderr
            config['document'].pop(key)
        # Mixed formats use the same batch pipeline, preserving document order.
        cfg.write_text(json.dumps(config))
        jobs = {'jobs': [{'input': str(root/'scan.tif'), 'output': str(root/'batch-tiff'), 'priority': 1},
                         {'input': str(root/'cells.csv'), 'output': str(root/'batch-csv'), 'priority': 2}]}
        (root/'jobs.json').write_text(json.dumps(jobs))
        subprocess.run([binary, '--config', str(cfg), '--batch', str(root/'jobs.json')], check=True, timeout=150)
        assert len(json.loads((root/'batch-tiff/result.json').read_text())['pages']) == 2
        assert (root/'batch-csv/result.md').is_file()
    print('PASS: extended formats, rendered content, multi-page order, bounds, converter failures and mixed batch')


if __name__ == '__main__':
    p = argparse.ArgumentParser()
    p.add_argument('binary'); p.add_argument('--external', action='store_true'); p.add_argument('--ofd-converter')
    args = p.parse_args()
    run(str(Path(args.binary).resolve()), args.external, args.ofd_converter)
